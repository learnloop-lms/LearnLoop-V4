import random
import io
import csv
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.database import get_db
from app import models, schemas, auth
from app.services.ai_engine import ai_engine
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/quizzes", tags=["Quizzes & Lifelines"])

@router.get("/template/download")
def download_quiz_template(current_user = Depends(auth.get_current_user)):
    """
    Generates and downloads a clean, pre-formatted Excel template (.xlsx)
    for teachers to bulk upload quiz questions.
    """
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Questions"

    headers = [
        "question_text", "option_a", "option_b", "option_c", "option_d",
        "correct_option", "cognitive_level", "blooms_level", "target_pt_ls",
        "question_type", "hint_text", "explanation"
    ]
    ws.append(headers)

    # Style header row
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Add sample template rows covering all question formats
    sample_rows = [
        [
            "What is the average time complexity of searching an element in a balanced Binary Search Tree?",
            "O(n)", "O(log n)", "O(n log n)", "O(1)",
            "B", "medium", "Understand", "Visionaries", "mcq",
            "Think of divide and conquer halving search space at each step",
            "A balanced BST halves the search space at each tree level leading to O(log n)."
        ],
        [
            "In Python, tuples are immutable data structures that cannot be modified after creation.",
            "True", "False", "", "",
            "A", "easy", "Remember", "universal", "true_false",
            "Tuples cannot be altered once defined",
            "Tuples are immutable in Python, whereas lists are mutable."
        ],
        [
            "The Java virtual machine converts compiled Java bytecode into _______ code during runtime execution.",
            "machine", "", "", "",
            "A", "medium", "Understand", "universal", "fill_blank",
            "Native hardware executable instructions",
            "The JIT compiler within JVM translates bytecode into native machine code."
        ],
        [
            "Match the programming language with its primary paradigm:\nList - I\tList - II\n(a) Java\t(i) Object Oriented\n(b) Python\t(ii) Multi-paradigm Interpreted\n(c) Prolog\t(iii) Logic Programming\n(d) C\t(iv) Procedural",
            "(a)-(i), (b)-(ii), (c)-(iii), (d)-(iv)",
            "(a)-(ii), (b)-(i), (c)-(iv), (d)-(iii)",
            "(a)-(iv), (b)-(iii), (c)-(i), (d)-(ii)",
            "(a)-(i), (b)-(iii), (c)-(ii), (d)-(iv)",
            "A", "hard", "Analyze", "universal", "match",
            "Consider Prolog's logic foundation and Java's class structure",
            "Java is Object Oriented, Python is Multi-paradigm, Prolog is Logic, and C is Procedural."
        ],
        [
            "What will be the output of the following Java snippet?\n\npublic class Test {\n    public static void main(String[] args) {\n        int a = 5;\n        System.out.println(a++ + ++a);\n    }\n}",
            "10", "11", "12", "13",
            "C", "hard", "Apply", "universal", "mcq",
            "x++ uses the current value then increments, while ++x increments first",
            "a++ evaluates to 5 and sets a=6. ++a increments a to 7 and evaluates to 7. Result: 5 + 7 = 12."
        ]
    ]

    for row in sample_rows:
        ws.append(row)

    # Adjust column widths automatically
    col_widths = {
        "A": 50, "B": 24, "C": 24, "D": 24, "E": 24,
        "F": 16, "G": 18, "H": 18, "I": 18, "J": 18,
        "K": 35, "L": 40
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers_response = {
        "Content-Disposition": "attachment; filename=learn_loop_quiz_template.xlsx"
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_response
    )

@router.post("/bulk-upload")
async def bulk_upload_quiz_questions(
    module_id: str = Form(...),
    file: UploadFile = File(...),
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    """
    Parses an uploaded Excel (.xlsx) or CSV file and inserts all questions into the database.
    """
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")

    module = db.query(models.Module).filter(models.Module.id == module_id).first()
    if not module:
        raise HTTPException(status_code=404, detail="Module not found.")

    filename = (file.filename or "").lower()
    contents = await file.read()
    if not contents:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    rows_data = []

    if filename.endswith(".xlsx") or filename.endswith(".xlsm"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                raise HTTPException(status_code=400, detail="Excel file must contain a header row and at least one question.")
            
            raw_headers = [str(c).strip().lower() if c else "" for c in rows[0]]
            
            # Map column indices
            def get_col(name, alt_names=None):
                if name in raw_headers:
                    return raw_headers.index(name)
                if alt_names:
                    for alt in alt_names:
                        if alt in raw_headers:
                            return raw_headers.index(alt)
                return -1

            q_idx = get_col("question_text", ["question", "question text", "q_text"])
            a_idx = get_col("option_a", ["option a", "opt_a", "a"])
            b_idx = get_col("option_b", ["option b", "opt_b", "b"])
            c_idx = get_col("option_c", ["option c", "opt_c", "c"])
            d_idx = get_col("option_d", ["option d", "opt_d", "d"])
            ans_idx = get_col("correct_option", ["correct", "answer", "correct answer", "correct_answer"])
            cog_idx = get_col("cognitive_level", ["difficulty", "complexity", "level"])
            bloom_idx = get_col("blooms_level", ["bloom", "blooms", "bloom_level"])
            pt_idx = get_col("target_pt_ls", ["pt_ls", "archetype", "target"])
            type_idx = get_col("question_type", ["type", "q_type"])
            hint_idx = get_col("hint_text", ["hint"])
            exp_idx = get_col("explanation", ["exp", "solution"])

            if q_idx == -1 or ans_idx == -1:
                raise HTTPException(status_code=400, detail="Missing required columns: question_text and correct_option.")

            for r_idx, r in enumerate(rows[1:], start=2):
                if not r or not any(r):
                    continue
                q_text = str(r[q_idx]).strip() if q_idx < len(r) and r[q_idx] is not None else ""
                if not q_text:
                    continue
                
                q_type = str(r[type_idx]).strip().lower() if type_idx != -1 and type_idx < len(r) and r[type_idx] else ""
                opt_a = str(r[a_idx]).strip() if a_idx != -1 and a_idx < len(r) and r[a_idx] is not None else ""
                opt_b = str(r[b_idx]).strip() if b_idx != -1 and b_idx < len(r) and r[b_idx] is not None else ""
                opt_c = str(r[c_idx]).strip() if c_idx != -1 and c_idx < len(r) and r[c_idx] is not None else ""
                opt_d = str(r[d_idx]).strip() if d_idx != -1 and d_idx < len(r) and r[d_idx] is not None else ""
                raw_ans = str(r[ans_idx]).strip() if ans_idx < len(r) and r[ans_idx] is not None else "A"

                rows_data.append({
                    "question_text": q_text,
                    "question_type": q_type,
                    "option_a": opt_a,
                    "option_b": opt_b,
                    "option_c": opt_c,
                    "option_d": opt_d,
                    "correct_option": raw_ans,
                    "cognitive_level": str(r[cog_idx]).strip().lower() if cog_idx != -1 and cog_idx < len(r) and r[cog_idx] else "medium",
                    "blooms_level": str(r[bloom_idx]).strip() if bloom_idx != -1 and bloom_idx < len(r) and r[bloom_idx] else "Understand",
                    "target_pt_ls": str(r[pt_idx]).strip() if pt_idx != -1 and pt_idx < len(r) and r[pt_idx] else "universal",
                    "hint_text": str(r[hint_idx]).strip() if hint_idx != -1 and hint_idx < len(r) and r[hint_idx] else None,
                    "explanation": str(r[exp_idx]).strip() if exp_idx != -1 and exp_idx < len(r) and r[exp_idx] else None,
                })
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    elif filename.endswith(".csv"):
        try:
            text = contents.decode("utf-8", errors="ignore")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                clean_row = {k.strip().lower(): v for k, v in row.items() if k}
                q_text = clean_row.get("question_text") or clean_row.get("question") or ""
                if not q_text.strip():
                    continue
                rows_data.append({
                    "question_text": q_text.strip(),
                    "question_type": (clean_row.get("question_type") or clean_row.get("type") or "").strip().lower(),
                    "option_a": (clean_row.get("option_a") or clean_row.get("a") or "").strip(),
                    "option_b": (clean_row.get("option_b") or clean_row.get("b") or "").strip(),
                    "option_c": (clean_row.get("option_c") or clean_row.get("c") or "").strip(),
                    "option_d": (clean_row.get("option_d") or clean_row.get("d") or "").strip(),
                    "correct_option": (clean_row.get("correct_option") or clean_row.get("correct") or "A").strip(),
                    "cognitive_level": (clean_row.get("cognitive_level") or clean_row.get("difficulty") or "medium").strip().lower(),
                    "blooms_level": (clean_row.get("blooms_level") or clean_row.get("bloom") or "Understand").strip(),
                    "target_pt_ls": (clean_row.get("target_pt_ls") or clean_row.get("pt_ls") or "universal").strip(),
                    "hint_text": (clean_row.get("hint_text") or clean_row.get("hint") or "").strip() or None,
                    "explanation": (clean_row.get("explanation") or clean_row.get("exp") or "").strip() or None,
                })
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to parse CSV file: {str(e)}")
    else:
        raise HTTPException(status_code=400, detail="Only .xlsx (Excel) and .csv files are supported.")

    if not rows_data:
        raise HTTPException(status_code=400, detail="No valid questions found in file.")

    created = []
    for item in rows_data:
        q_text = item["question_text"]
        q_type = item.get("question_type", "").lower()
        opt_a = item["option_a"]
        opt_b = item["option_b"]
        opt_c = item["option_c"]
        opt_d = item["option_d"]
        raw_corr = item["correct_option"].strip()

        # Check if question is True/False
        is_tf = (
            q_type in ["true_false", "tf", "boolean"] or
            (opt_a.lower() == "true" and opt_b.lower() == "false") or
            (opt_a.lower() == "false" and opt_b.lower() == "true") or
            "[true/false]" in q_text.lower() or
            q_text.lower().startswith("true or false")
        )

        # Check if question is Fill in the Blank
        is_fib = (
            q_type in ["fill_blank", "fill_in_the_blank", "fib", "blank"] or
            "___" in q_text or
            "[blank]" in q_text.lower() or
            ("fill in the blank" in q_text.lower() and not opt_b)
        )

        if is_tf:
            if not opt_a:
                opt_a = "True"
            if not opt_b:
                opt_b = "False"
            opt_c = ""
            opt_d = ""
            if raw_corr.upper() in ["A", "TRUE", "T", "1"]:
                corr = "A"
            else:
                corr = "B"
        elif is_fib:
            # In Fill in the blank: answer string is in option_a
            if not opt_a and raw_corr and raw_corr.upper() not in ["A", "B", "C", "D"]:
                opt_a = raw_corr
            opt_b = ""
            opt_c = ""
            opt_d = ""
            corr = "A"
        else:
            # Standard MCQ / Match
            corr = raw_corr[:1].upper() if raw_corr else "A"
            if corr not in ["A", "B", "C", "D"]:
                corr = "A"
            if not opt_a:
                opt_a = "Option A"
            if not opt_b:
                opt_b = "Option B"
            # Do NOT inject dummy "Option C" or "Option D" if they were not provided

        # Validate cognitive level
        cog = item["cognitive_level"].lower()
        if cog in ["easy", "slow", "basic"]:
            cog = "slow"
        elif cog in ["hard", "fast", "advanced"]:
            cog = "fast"
        else:
            cog = "average"

        q = models.QuizQuestion(
            module_id=module_id,
            cognitive_level=cog,
            target_pt_ls=item["target_pt_ls"] or "Universal",
            blooms_level=item["blooms_level"] or "Understand",
            question_text=q_text,
            option_a=opt_a,
            option_b=opt_b,
            option_c=opt_c,
            option_d=opt_d,
            correct_option=corr,
            hint_text=item["hint_text"],
            explanation=item["explanation"]
        )
        db.add(q)
        created.append(q)

    db.commit()
    return {
        "message": f"Successfully imported {len(created)} quiz questions from Excel!",
        "count": len(created)
    }

@router.post("/questions")
def add_quiz_question(
    req: schemas.QuizQuestionCreate,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")

    q = models.QuizQuestion(
        module_id=req.module_id,
        cognitive_level=req.cognitive_level.lower(),
        target_pt_ls=req.target_pt_ls,
        blooms_level=req.blooms_level,
        question_text=req.question_text,
        option_a=req.option_a,
        option_b=req.option_b,
        option_c=req.option_c,
        option_d=req.option_d,
        correct_option=req.correct_option.upper().strip(),
        hint_text=req.hint_text,
        explanation=req.explanation
    )
    db.add(q)
    db.commit()
    db.refresh(q)
    return {"message": "Question added successfully", "id": q.id}

@router.get("/module/{module_id}/all")
def get_all_module_quiz_questions(
    module_id: str,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")
    return db.query(models.QuizQuestion).filter(models.QuizQuestion.module_id == module_id).all()

@router.put("/questions/{question_id}")
def update_quiz_question(
    question_id: str,
    req: schemas.QuizQuestionUpdate,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")

    q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Quiz question not found")

    if req.question_text is not None:
        q.question_text = req.question_text.strip()
    if req.cognitive_level is not None:
        q.cognitive_level = req.cognitive_level.lower().strip()
    if req.target_pt_ls is not None:
        q.target_pt_ls = req.target_pt_ls.strip()
    if req.blooms_level is not None:
        q.blooms_level = req.blooms_level.strip()
    if req.option_a is not None:
        q.option_a = req.option_a
    if req.option_b is not None:
        q.option_b = req.option_b
    if req.option_c is not None:
        q.option_c = req.option_c
    if req.option_d is not None:
        q.option_d = req.option_d
    if req.correct_option is not None:
        q.correct_option = req.correct_option.strip()
    if req.hint_text is not None:
        q.hint_text = req.hint_text
    if req.explanation is not None:
        q.explanation = req.explanation

    db.commit()
    db.refresh(q)
    return {"message": "Quiz question updated successfully", "id": q.id}


@router.delete("/questions/{question_id}")
def delete_quiz_question(
    question_id: str,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "teacher":
        raise HTTPException(status_code=403, detail="Teachers only.")

    q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Quiz question not found")

    db.delete(q)
    db.commit()
    return {"message": "Quiz question deleted successfully", "id": question_id}

@router.get("/module/{module_id}")
def get_adaptive_quiz(
    module_id: str,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "student":
        raise HTTPException(status_code=403, detail="Students only.")

    module = db.query(models.Module).filter(models.Module.id == module_id).first()
    if not module:
        raise HTTPException(status_code=404, detail="Module not found.")

    student_cog = (current_user.cognitive_level or "average").lower()
    student_ptls = (current_user.pt_ls or "Universal").strip()

    # Dynamic breakdown of Q_e (Easy), Q_m (Medium), Q_h (Hard)
    if module.module_no == 1 and student_cog == "average":
        q_easy_req, q_med_req, q_hard_req = 5, 3, 2
    elif student_cog == "slow":
        q_easy_req, q_med_req, q_hard_req = 8, 2, 0
    elif student_cog == "average":
        q_easy_req, q_med_req, q_hard_req = 4, 4, 2
    else: # fast
        q_easy_req, q_med_req, q_hard_req = 2, 4, 4

    # Player-Type matching helpers
    def is_pt_specific(q_pt: str, s_pt: str) -> bool:
        if not q_pt or not q_pt.strip():
            return False
        qp = q_pt.strip().lower()
        sp = s_pt.strip().lower()
        if qp in ["universal", "all", "general", "any", "common"]:
            return False
        return qp.startswith(sp[:5]) or sp.startswith(qp[:5])

    def pt_matches(q_pt: str, s_pt: str) -> bool:
        if not q_pt or not q_pt.strip():
            return True
        qp = q_pt.strip().lower()
        sp = s_pt.strip().lower()
        if qp in ["universal", "all", "general", "any", "common"]:
            return True
        return qp.startswith(sp[:5]) or sp.startswith(qp[:5])

    # Fetch all questions for this module
    all_questions = db.query(models.QuizQuestion).filter(models.QuizQuestion.module_id == module_id).all()

    # Eligible questions = Archetype-matched questions + Universal questions (always shared with all students)
    eligible_pool = [q for q in all_questions if pt_matches(q.target_pt_ls, student_ptls)]
    if not eligible_pool:
        eligible_pool = all_questions

    # Cognitive pool grouping
    def get_cog_tier(cog: str) -> str:
        c = (cog or "").strip().lower()
        if c in ["slow", "easy", "basic"]:
            return "slow"
        elif c in ["fast", "hard", "advanced"]:
            return "fast"
        return "average"

    easy_pool = [q for q in eligible_pool if get_cog_tier(q.cognitive_level) == "slow"]
    med_pool = [q for q in eligible_pool if get_cog_tier(q.cognitive_level) == "average"]
    hard_pool = [q for q in eligible_pool if get_cog_tier(q.cognitive_level) == "fast"]

    selected = []
    selected += random.sample(easy_pool, min(len(easy_pool), q_easy_req))
    selected += random.sample(med_pool, min(len(med_pool), q_med_req))
    selected += random.sample(hard_pool, min(len(hard_pool), q_hard_req))

    # Fill up to 10 from remaining eligible questions
    remaining = [q for q in eligible_pool if q not in selected]
    random.shuffle(remaining)
    while len(selected) < min(10, len(eligible_pool)) and remaining:
        selected.append(remaining.pop(0))

    random.shuffle(selected)

    return {
        "module_id": module_id,
        "concept_name": module.concept_name,
        "student_cognitive_level": student_cog,
        "student_player_type": student_ptls,
        "total_questions": len(selected),
        "questions": [
            {
                "id": q.id,
                "cognitive_level": q.cognitive_level,
                "blooms_level": q.blooms_level,
                "target_pt_ls": q.target_pt_ls,
                "question_text": q.question_text,
                "option_a": q.option_a,
                "option_b": q.option_b,
                "option_c": q.option_c,
                "option_d": q.option_d,
                "hint_available": bool(q.hint_text)
            }
            for q in selected
        ],
        "lifelines_remaining": {
            "hearts": current_user.hearts,
            "coins": current_user.total_coins,
            "free_hints": current_user.free_hints,
            "free_extra_time": current_user.free_extra_time,
            "free_swaps": current_user.free_swaps,
            "free_double_dips": current_user.free_double_dips,
            "free_fifty_fifty": current_user.free_fifty_fifty,
            "high_achiever_bonus_active": current_user.high_achiever_bonus_active
        }
    }

@router.post("/lifelines/use")
def use_lifeline(
    req: schemas.LifelineRequest,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "student":
        raise HTTPException(status_code=403, detail="Students only.")

    l_type = req.lifeline_type.lower()
    bonus_active = current_user.high_achiever_bonus_active

    if l_type == "50_50":
        if not bonus_active:
            if current_user.free_fifty_fifty > 0:
                current_user.free_fifty_fifty -= 1
            else:
                if current_user.total_coins < 100:
                    raise HTTPException(status_code=400, detail="Not enough coins for 50:50 (Costs 100 coins).")
                current_user.total_coins -= 100

        # Find question and eliminate two wrong options
        q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == req.question_id).first()
        if not q:
            raise HTTPException(status_code=404, detail="Question not found.")

        all_options = ["A", "B", "C", "D"]
        correct = q.correct_option.upper()
        wrong_options = [opt for opt in all_options if opt != correct]
        eliminated = random.sample(wrong_options, 2)
        db.commit()

        return {
            "success": True,
            "eliminated_options": eliminated,
            "message": f"Eliminated options: {', '.join(eliminated)}"
        }

    elif l_type == "double_dip":
        if not bonus_active:
            if current_user.free_double_dips > 0:
                current_user.free_double_dips -= 1
            else:
                if current_user.total_coins < 100:
                    raise HTTPException(status_code=400, detail="Not enough coins for Double Dip (Costs 100 coins).")
                current_user.total_coins -= 100
        db.commit()
        return {"success": True, "message": "Double Dip activated! You can select 2 choices for this question."}

    elif l_type == "hint":
        if not bonus_active:
            if current_user.free_hints > 0:
                current_user.free_hints -= 1
            else:
                if current_user.total_coins < 10:
                    raise HTTPException(status_code=400, detail="Not enough coins for Hint (Costs 10 coins).")
                current_user.total_coins -= 10

        q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == req.question_id).first()
        hint = q.hint_text if q else "Think carefully about the key concepts of this module."
        db.commit()
        return {"success": True, "hint": hint}

    elif l_type == "swap":
        if not bonus_active:
            if current_user.free_swaps > 0:
                current_user.free_swaps -= 1
            else:
                if current_user.hearts <= 1:
                    raise HTTPException(status_code=400, detail="Not enough lives to swap question (Costs 1 life).")
                current_user.hearts -= 1
                current_user.last_heart_loss_at = datetime.utcnow()

        # Find replacement question
        orig_q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == req.question_id).first()
        replacement = db.query(models.QuizQuestion).filter(
            models.QuizQuestion.module_id == req.module_id,
            models.QuizQuestion.id != req.question_id,
            models.QuizQuestion.cognitive_level == orig_q.cognitive_level if orig_q else True
        ).first()

        db.commit()
        if not replacement:
            return {"success": False, "message": "No other question available in pool to swap."}

        return {
            "success": True,
            "new_question": {
                "id": replacement.id,
                "cognitive_level": replacement.cognitive_level,
                "blooms_level": replacement.blooms_level,
                "question_text": replacement.question_text,
                "option_a": replacement.option_a,
                "option_b": replacement.option_b,
                "option_c": replacement.option_c,
                "option_d": replacement.option_d,
                "hint_available": bool(replacement.hint_text)
            }
        }

    elif l_type == "extra_time":
        if not bonus_active:
            if current_user.free_extra_time > 0:
                current_user.free_extra_time -= 1
            else:
                if current_user.total_coins < 50:
                    raise HTTPException(status_code=400, detail="Not enough coins for Extra Time (Costs 50 coins).")
                current_user.total_coins -= 50
        db.commit()
        return {"success": True, "additional_seconds": 60, "message": "+1 Minute Added!"}

    raise HTTPException(status_code=400, detail="Unknown lifeline type.")

@router.post("/submit", response_model=schemas.QuizSubmissionOut)
def submit_quiz(
    req: schemas.QuizSubmissionRequest,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    if getattr(current_user, "current_role", "") != "student":
        raise HTTPException(status_code=403, detail="Students only.")

    module = db.query(models.Module).filter(models.Module.id == req.module_id).first()
    if not module:
        raise HTTPException(status_code=404, detail="Module not found.")

    course = module.course
    passing_threshold = course.passing_threshold_p if course else 5.0

    score = 0
    mistakes_by_cat = {"slow": 0, "average": 0, "fast": 0}
    q_counts = {"slow": 0, "average": 0, "fast": 0}
    explanation_map = {}
    is_visionary = current_user.pt_ls == "Visionaries"
    detailed_responses = []

    for ans in req.answers:
        q = db.query(models.QuizQuestion).filter(models.QuizQuestion.id == ans.question_id).first()
        if not q:
            continue

        q_cat = q.cognitive_level.lower()
        if q_cat in ["easy", "basic"]:
            q_cat = "slow"
        elif q_cat in ["hard", "advanced"]:
            q_cat = "fast"
        else:
            q_cat = "average"

        q_counts[q_cat] = q_counts.get(q_cat, 0) + 1

        opt_a_clean = (q.option_a or "").strip()
        opt_b_clean = (q.option_b or "").strip()
        opt_c_clean = (q.option_c or "").strip()
        opt_d_clean = (q.option_d or "").strip()

        # Identify True/False
        is_tf = (
            (opt_a_clean.lower() == "true" and opt_b_clean.lower() == "false") or
            (opt_a_clean.lower() == "false" and opt_b_clean.lower() == "true") or
            "[true/false]" in q.question_text.lower() or
            q.question_text.lower().startswith("true or false")
        )

        # Identify Fill in the Blank
        is_fib = (
            "___" in q.question_text or
            "[blank]" in q.question_text.lower() or
            ("fill in the blank" in q.question_text.lower() and not opt_b_clean) or
            (not opt_b_clean and not opt_c_clean and not opt_d_clean)
        )

        student_ans_raw = (ans.chosen_option or "").strip()
        is_correct = False
        correct_answer_text = ""

        if is_fib:
            # Case-insensitive comparison with all lowercase
            expected_answer = (opt_a_clean or q.correct_option or "").strip().lower()
            student_clean = student_ans_raw.strip().lower()
            is_correct = bool(student_clean and expected_answer and (student_clean == expected_answer))
            correct_answer_text = opt_a_clean or q.correct_option
        elif is_tf:
            corr_letter = (q.correct_option or "A").strip().upper()
            if opt_a_clean.lower() == "true":
                expected_tf = "TRUE" if corr_letter in ["A", "TRUE", "T"] else "FALSE"
            else:
                expected_tf = "FALSE" if corr_letter in ["A", "FALSE"] else "TRUE"

            student_tf = "TRUE" if student_ans_raw.upper() in ["A", "TRUE", "T"] else "FALSE"
            is_correct = (student_tf == expected_tf)
            correct_answer_text = expected_tf.capitalize()
        else:
            # Standard MCQ / Match
            corr_letter = (q.correct_option or "A").strip().upper()[:1]
            student_letter = student_ans_raw.strip().upper()[:1]
            is_correct = (student_letter == corr_letter)
            opt_map = {"A": opt_a_clean, "B": opt_b_clean, "C": opt_c_clean, "D": opt_d_clean}
            correct_answer_text = f"{corr_letter}: {opt_map.get(corr_letter, opt_a_clean)}"

        if is_correct:
            score += 1
        else:
            mistakes_by_cat[q_cat] = mistakes_by_cat.get(q_cat, 0) + 1

        if is_visionary and q.explanation:
            explanation_map[q.id] = q.explanation

        detailed_responses.append({
            "question_id": q.id,
            "question_text": q.question_text,
            "option_a": opt_a_clean,
            "option_b": opt_b_clean,
            "option_c": opt_c_clean,
            "option_d": opt_d_clean,
            "chosen_option": ans.chosen_option,
            "student_answer": student_ans_raw,
            "correct_option": q.correct_option,
            "correct_answer": correct_answer_text,
            "is_correct": is_correct,
            "explanation": q.explanation or q.hint_text,
            "is_fill_blank": is_fib,
            "is_tf": is_tf,
            "lifelines_used": getattr(ans, "lifelines_used", []) or []
        })

    total_q = len(req.answers) if req.answers else 1
    accuracy = score / total_q

    # Determine HIT (mistake category with highest incorrect count)
    hit_cat = max(mistakes_by_cat, key=mistakes_by_cat.get) if any(mistakes_by_cat.values()) else "none"
    hit_encoding_map = {"slow": 1, "average": 2, "fast": 3, "none": 0}
    hit_encoded = hit_encoding_map.get(hit_cat, 0)

    # -------------------------------------------------------------------------
    # ANN Model 2: Evaluate Quiz Cognitive Level (STRICTLY INTERNAL coq)
    # -------------------------------------------------------------------------
    coq_predicted = ai_engine.predict_quiz_cognitive(
        msq=float(score),
        ttq_sec=req.ttq_seconds,
        accuracy=accuracy,
        hit_encoded=hit_encoded
    )

    # Record submission with full detailed breakdown
    submission = models.QuizSubmission(
        student_id=current_user.id,
        module_id=req.module_id,
        msq_score=float(score),
        ttq_seconds=req.ttq_seconds,
        mistake_category_hit=hit_cat,
        q_easy_count=q_counts.get("slow", 0),
        q_med_count=q_counts.get("average", 0),
        q_hard_count=q_counts.get("fast", 0),
        coq_predicted=coq_predicted,
        responses_detail=detailed_responses
    )
    db.add(submission)

    passed = score >= passing_threshold

    if not passed:
        # Algorithm 5: Downgrade cognitive level by 1 position
        if current_user.cognitive_level == "fast":
            current_user.cognitive_level = "average"
        elif current_user.cognitive_level == "average":
            current_user.cognitive_level = "slow"
        
        # Deduct 2 lives if failing easy level
        if hit_cat == "slow" or current_user.cognitive_level == "slow":
            current_user.hearts = max(0, current_user.hearts - 2)
            current_user.last_heart_loss_at = datetime.utcnow()

        next_step = "downgraded_retry"
    else:
        next_step = "proceed_to_task"

    db.commit()

    return schemas.QuizSubmissionOut(
        id=submission.id,
        msq_score=float(score),
        total_questions=total_q,
        accuracy=accuracy,
        mistake_category_hit=hit_cat,
        passed=passed,
        next_step=next_step,
        explanation_map=explanation_map if is_visionary else None
    )

@router.get("/module/{module_id}/my-submission")
def get_my_quiz_submission(
    module_id: str,
    current_user = Depends(auth.get_current_user),
    db: Session = Depends(get_db)
):
    """
    Returns the student's completed quiz submission for this module,
    including each question, chosen answer, correct answer, and correctness.
    """
    if getattr(current_user, "current_role", "") != "student":
        raise HTTPException(status_code=403, detail="Students only.")

    sub = db.query(models.QuizSubmission).filter(
        models.QuizSubmission.module_id == module_id,
        models.QuizSubmission.student_id == current_user.id
    ).order_by(models.QuizSubmission.submitted_at.desc()).first()

    if not sub:
        raise HTTPException(status_code=404, detail="No quiz submission found for this module.")

    module = db.query(models.Module).filter(models.Module.id == module_id).first()
    course = module.course if module else None
    passing_threshold = course.passing_threshold_p if course else 5.0

    return {
        "id": sub.id,
        "module_id": sub.module_id,
        "concept_name": module.concept_name if module else "Quiz",
        "msq_score": sub.msq_score,
        "total_questions": len(sub.responses_detail) if sub.responses_detail else 10,
        "passing_threshold": passing_threshold,
        "passed": sub.msq_score >= passing_threshold,
        "submitted_at": sub.submitted_at,
        "responses_detail": sub.responses_detail or []
    }
